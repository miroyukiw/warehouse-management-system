import tkinter as tk
from tkinter import ttk, messagebox
import sqlite3
import pandas as pd
from openpyxl import load_workbook
import datetime
import os
import subprocess
import shutil

import win32com.client as win32  # 关键：Excel COM，用来保留页眉图片

DB_FILE = 'warehouse.db'

def connect_db():
    return sqlite3.connect(DB_FILE)

def get_distinct_outbound_values(field):
    conn = connect_db()
    c = conn.cursor()
    c.execute(f"SELECT DISTINCT {field} FROM 出库表 WHERE {field} IS NOT NULL")
    values = [r[0] for r in c.fetchall() if r[0] is not None]
    conn.close()
    return values

def get_batch_data_for(material, spec):
    conn = connect_db()
    c = conn.cursor()
    c.execute("""
        SELECT 序号, 批次或卷号, 剩余数量, 备注
        FROM 批次库存表
        WHERE 物资名称=? AND 规格型号=? AND 剩余数量>0
        ORDER BY 序号 ASC
    """, (material, spec))
    rows = c.fetchall()
    conn.close()
    return rows

# =============== EditableTable 类 ===============
class EditableTable(ttk.Treeview):
    def __init__(self, parent, table_name, columns, **kwargs):
        # 入库、出库表 => 多选；库存表 => 单选
        selectmode = 'extended' if table_name in ("入库表", "出库表") else 'browse'
        super().__init__(parent, columns=columns, show='headings', selectmode=selectmode, **kwargs)
        self["show"] = "headings"
        self.table_name = table_name
        self.columns_list = list(columns)

        # 如果是“库存表”，最后加一列“展开”
        if self.table_name == "库存表":
            self.columns_list.append("展开")
            self["columns"] = tuple(self.columns_list)

        self.full_data = []
        self.sort_order = {}
        self.current_filter = ""
        self.expanded_rows = {}

        # 设置表头样式
        for col in self.columns_list:
            self.heading(col, text=col, anchor='center')
            self.column(col, anchor='center', width=120)

        # 绑定事件
        self.bind("<Button-3>", self.on_header_right_click)
        self.bind("<Double-1>", self.on_double_click)
        if self.table_name in ("入库表", "出库表"):
            self.bind("<Button-1>", self.on_header_left_click, add="+")
        if self.table_name == "库存表":
            self.bind("<Button-1>", self.on_click_cell, add="+")

        # 滚动条
        vsb = ttk.Scrollbar(parent, orient="vertical", command=self.yview)
        vsb.pack(side="right", fill="y")
        self.configure(yscrollcommand=vsb.set)
        hsb = ttk.Scrollbar(parent, orient="horizontal", command=self.xview)
        hsb.pack(side="bottom", fill="x")
        self.configure(xscrollcommand=hsb.set)

    def load_data(self, query_filter=""):
        # 先清空
        for item in self.get_children():
            self.delete(item)

        conn = connect_db()
        c = conn.cursor()

        # ---------------------------
        # 根据需求，入库表/出库表默认按日期倒序
        # ---------------------------
        if self.table_name == "入库表":
            sql = "SELECT rowid, * FROM 入库表 ORDER BY 入库日期 DESC"
        elif self.table_name == "出库表":
            sql = "SELECT rowid, * FROM 出库表 ORDER BY 出库日期 DESC"
        elif self.table_name == "库存表" and hasattr(self, 'hide_zero') and self.hide_zero:
            sql = "SELECT rowid, * FROM 库存表 WHERE 当前库存 > 0"
        else:
            sql = f"SELECT rowid, * FROM {self.table_name}"

        c.execute(sql)
        rows = c.fetchall()
        conn.close()
        self.full_data = rows

        for row in rows:
            rowid = row[0]
            # 将 None / "NULL" 转为空字符串
            vals = [("" if (v is None or str(v).upper()=="NULL") else v) for v in row[1:]]

            # 如果是库存表，最后一列看是否有批次库存
            if self.table_name == "库存表":
                has_batch = self.check_batch_exists(vals)
                plus_val = "＋" if has_batch else ""
                vals.append(plus_val)

            # 如果用户输入了搜索关键字
            if query_filter:
                # 搜索范围：物资名称、规格型号、厂家
                indices = []
                for col in ["物资名称", "规格型号", "厂家"]:
                    if col in self.columns_list:
                        indices.append(self.columns_list.index(col))
                row_text = " ".join(str(vals[i]) for i in indices)
                if query_filter.lower() not in row_text.lower():
                    continue

            self.insert("", "end", iid=rowid, values=vals)

    def filter_data(self, keyword):
        self.current_filter = keyword
        self.load_data(query_filter=keyword)

    def check_batch_exists(self, row_values):
        """判断表格行对应的物资名称、规格型号是否在批次库存表里仍有库存"""
        try:
            mat = row_values[1]
            spc = row_values[2]
        except:
            return False
        conn = connect_db()
        c = conn.cursor()
        c.execute("""SELECT COUNT(*) FROM 批次库存表
                     WHERE 物资名称=? AND 规格型号=? AND 剩余数量>0""", (mat, spc))
        cnt = c.fetchone()[0]
        conn.close()
        return (cnt > 0)

    def on_header_right_click(self, event):
        region = self.identify("region", event.x, event.y)
        if region != "heading":
            return
        colid = self.identify_column(event.x)
        col_idx = int(colid.replace("#", "")) - 1
        if col_idx < 0 or col_idx >= len(self.columns_list):
            return
        col_name = self.columns_list[col_idx]

        import datetime
        if col_name in ("入库日期", "出库日期"):
            menu = tk.Menu(self, tearoff=0)
            menu.add_command(label="升序", command=lambda: self.sort_by(col_name, ascending=True))
            menu.add_command(label="降序", command=lambda: self.sort_by(col_name, ascending=False))
            menu.tk_popup(event.x_root, event.y_root)
        else:
            menu = tk.Menu(self, tearoff=0)
            menu.add_command(label="隐藏此列", command=lambda: self.hide_column(col_name))
            menu.add_command(label="显示所有列", command=self.show_all_columns)
            menu.tk_popup(event.x_root, event.y_root)

    def sort_by(self, col_name, ascending=True):
        try:
            idx = self.columns_list.index(col_name)
        except:
            return
        import datetime

        def parse_date(val):
            if not val:
                return datetime.datetime.min
            for fmt in ("%Y/%m/%d", "%Y-%m-%d"):
                try:
                    return datetime.datetime.strptime(val, fmt)
                except:
                    pass
            return datetime.datetime.min

        def sort_key(row):
            val = row[1 + idx]
            if col_name in ("入库日期", "出库日期"):
                return parse_date(val)
            return val if val else ""

        self.full_data.sort(key=sort_key, reverse=not ascending)
        self.load_data(query_filter=self.current_filter)

    def hide_column(self, col_name):
        self.column(col_name, width=0, minwidth=0)

    def show_all_columns(self):
        for col in self.columns_list:
            self.column(col, width=120, minwidth=50)

    def on_header_left_click(self, event):
        pass

    def on_double_click(self, event):
        """双击处理：若是库存表的“展开”列 -> toggle_expand；否则编辑单元格"""
        region = self.identify("region", event.x, event.y)
        if region == "heading":
            colid = self.identify_column(event.x)
            col_idx = int(colid.replace("#", "")) - 1
            # 出库表：双击“项目名称”“项目编码”“领用人” 可弹出筛选
            if (self.table_name == "出库表" and
                self.columns_list[col_idx] in ("项目名称", "项目编码", "领用人")):
                self.show_filter_window(self.columns_list[col_idx])
            return

        colid = self.identify_column(event.x)
        col_idx = int(colid.replace("#", "")) - 1

        # 库存表“展开”列
        if self.table_name == "库存表" and col_idx == len(self.columns_list) - 1:
            row_id = self.identify_row(event.y)
            self.toggle_expand(row_id)
            return

        # 进入编辑模式
        item = self.identify_row(event.y)
        if not item:
            return
        x, y, w, h = self.bbox(item, colid)
        entry = ttk.Entry(self)
        entry.place(x=x, y=y, width=w, height=h)
        entry.insert(0, self.set(item, self.columns_list[col_idx]))
        entry.focus()

        def on_save(e):
            new_val = entry.get()
            entry.destroy()
            self.set(item, self.columns_list[col_idx], new_val)
            conn = connect_db()
            c = conn.cursor()
            c.execute(f"UPDATE {self.table_name} SET {self.columns_list[col_idx]}=? WHERE rowid=?",
                      (new_val, item))
            conn.commit()
            conn.close()

        entry.bind("<Return>", on_save)
        entry.bind("<FocusOut>", lambda e: entry.destroy())

    def show_filter_window(self, col_name):
        """出库表：可对指定列值做筛选"""
        top = tk.Toplevel(self)
        top.title(f"筛选 {col_name}")
        top.geometry("1280x720")

        style = ttk.Style(top)
        style.theme_use("clam")

        search_frame = ttk.Frame(top)
        search_frame.pack(fill="x", padx=10, pady=10)
        ttk.Label(search_frame, text="搜索:").pack(side="left")
        search_var = tk.StringVar()
        search_entry = ttk.Entry(search_frame, textvariable=search_var, width=40)
        search_entry.pack(side="left", padx=5)

        def do_filter():
            term = search_var.get().strip().lower()
            listbox.delete(0, tk.END)
            for v in distinct_vals:
                if term in str(v).lower():
                    listbox.insert(tk.END, v)

        search_btn = ttk.Button(search_frame, text="筛选", command=do_filter)
        search_btn.pack(side="left")

        frame = ttk.Frame(top)
        frame.pack(fill="both", expand=True, padx=10, pady=10)
        vscroll = ttk.Scrollbar(frame, orient="vertical")
        hscroll = ttk.Scrollbar(frame, orient="horizontal")
        listbox = tk.Listbox(frame, yscrollcommand=vscroll.set, xscrollcommand=hscroll.set, font=("微软雅黑",12))
        vscroll.config(command=listbox.yview)
        hscroll.config(command=listbox.xview)
        vscroll.pack(side="right", fill="y")
        hscroll.pack(side="bottom", fill="x")
        listbox.pack(fill="both", expand=True)

        distinct_vals = set()
        try:
            idx = self.columns_list.index(col_name)
        except:
            idx = None
        if idx is not None:
            for row in self.full_data:
                val = row[1+idx]
                if val not in (None, "", "NULL"):
                    distinct_vals.add(val)
        distinct_vals = sorted(list(distinct_vals))
        for v in distinct_vals:
            listbox.insert(tk.END, v)

        def on_select(e):
            sel = listbox.curselection()
            if sel:
                value = listbox.get(sel[0])
                self.filter_by_column(col_name, value)
            top.destroy()

        listbox.bind("<Double-Button-1>", on_select)

    def filter_by_column(self, col_name, value):
        """只显示该列值=某value的行"""
        try:
            idx = self.columns_list.index(col_name)
        except:
            return
        filtered = [row for row in self.full_data if row[1 + idx] == value]
        self.delete(*self.get_children())
        for row in filtered:
            rowid = row[0]
            vals = [("" if (v is None or str(v).upper()=="NULL") else v) for v in row[1:]]
            if self.table_name == "库存表":
                has_batch = self.check_batch_exists(vals)
                plus_val = "＋" if has_batch else ""
                vals.append(plus_val)
            self.insert("", "end", iid=rowid, values=vals)

    def on_click_cell(self, event):
        region = self.identify("region", event.x, event.y)
        if region != "cell":
            return
        item = self.identify_row(event.y)
        colid = self.identify_column(event.x)
        col_idx = int(colid.replace("#","")) - 1
        if not item or col_idx != len(self.columns_list) - 1:
            return
        self.toggle_expand(item)

    def toggle_expand(self, parent_id):
        if parent_id in self.expanded_rows:
            # 收起
            for child in self.expanded_rows[parent_id]:
                self.delete(child)
            del self.expanded_rows[parent_id]
            vals = list(self.item(parent_id,"values"))
            if self.check_batch_exists(vals):
                vals[-1] = "＋"
            else:
                vals[-1] = ""
            self.item(parent_id, values=vals)
        else:
            # 展开
            pvals = self.item(parent_id,"values")
            if len(pvals) < 3:
                return
            mat = pvals[1]
            spc = pvals[2]
            childrows = get_batch_data_for(mat, spc)
            if not childrows:
                return
            child_ids = []
            idx = self.index(parent_id) + 1
            for (bid, bc, remain, remark) in childrows:
                child_vals = ["" for _ in range(len(self.columns_list))]
                child_vals[1] = f"批次: {bc}"
                child_vals[2] = f"剩余: {remain}"
                child_vals[3] = f"备注: {remark}"
                cid = f"{parent_id}-child-{bid}"
                self.insert("", idx, iid=cid, values=child_vals)
                child_ids.append(cid)
                idx += 1
            self.expanded_rows[parent_id] = child_ids
            vals = list(pvals)
            vals[-1] = "－"
            self.item(parent_id, values=vals)

# =============== InventoryTable 类 ===============
class InventoryTable(EditableTable):
    def __init__(self, parent, table_name, columns, **kwargs):
        super().__init__(parent, table_name, columns, **kwargs)
        # 默认不隐藏零库存，避免库存表空白
        self.hide_zero = False

# =============== InventoryApp 主界面 ===============
class InventoryApp:
    def __init__(self, root):
        self.root = root
        self.root.title("仓库管理界面 - 修复库存表 & 周报带库存")
        self.root.geometry("1200x750")

        style = ttk.Style(root)
        style.theme_use("clam")
        style.configure("Treeview", rowheight=25, font=('微软雅黑',10))
        style.configure("Treeview.Heading", background="#f0f0f0", font=('微软雅黑',11,'bold'))

        top_frame = ttk.Frame(root)
        top_frame.pack(fill="x", padx=5, pady=5)

        # “生成项目周报”按钮
        weekly_btn = ttk.Button(top_frame, text="生成项目周报", command=self.generate_weekly_report)
        weekly_btn.pack(side="left", padx=5)

        # 搜索区
        search_frame = ttk.Frame(top_frame)
        search_frame.pack(side="right")
        ttk.Label(search_frame, text="搜索:").pack(side="left")
        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(search_frame, textvariable=self.search_var, width=20)
        search_entry.pack(side="left", padx=5)
        search_btn = ttk.Button(search_frame, text="搜索", command=self.refresh_all)
        search_btn.pack(side="left")

        # Notebook：入库表 / 出库表 / 库存表
        notebook = ttk.Notebook(root)
        notebook.pack(fill="both", expand=True)
        self.notebook = notebook

        self.tables = {}

        # 列配置
        self.inbound_cols = [
            '序号','入库日期','物资名称','规格型号','专业类别','单位','数量','批次或卷号',
            '厂家','领料人','验收人','备注','项目名称','项目编码','委托人','物资白单单号',
            '物资属性','明细账存货编码','物资黄单单号'
        ]
        outbound_cols = [
            '序号','出库日期','物资名称','规格型号','厂家','项目名称','项目编码','单位',
            '数量','批次或卷号','领用人','经办人','备注'
        ]
        inventory_cols = [
            '序号','物资名称','规格型号','厂家','专业类别','委托人','单位',
            '当前库存','最近入库日期','最近出库日期','存放位置','备注'
        ]

        # ——— 入库表分页 ———
        inbound_frame = ttk.Frame(notebook)
        inbound_frame.pack(fill="both", expand=True)
        tip_label = ttk.Label(inbound_frame, text="(多选后点下面按钮)")
        tip_label.pack(fill="x", pady=5)
        export_inbound_btn = ttk.Button(inbound_frame, text="生成入库表", command=self.export_inbound)
        export_inbound_btn.pack(fill="x", pady=5)
        inbound_table_frame = ttk.Frame(inbound_frame)
        inbound_table_frame.pack(fill="both", expand=True)
        inbound_table = EditableTable(inbound_table_frame, "入库表", self.inbound_cols)
        inbound_table.pack(fill="both", expand=True)
        self.tables["入库表"] = inbound_table
        notebook.add(inbound_frame, text="入库表")

        # ——— 出库表分页 ———
        outbound_frame = ttk.Frame(notebook)
        outbound_frame.pack(fill="both", expand=True)
        export_outbound_btn = ttk.Button(outbound_frame, text="生成出库领料单", command=self.export_outbound)
        export_outbound_btn.pack(fill="x", pady=5)
        outbound_table_frame = ttk.Frame(outbound_frame)
        outbound_table_frame.pack(fill="both", expand=True)
        outbound_table = EditableTable(outbound_table_frame, "出库表", outbound_cols)
        outbound_table.pack(fill="both", expand=True)
        self.tables["出库表"] = outbound_table
        notebook.add(outbound_frame, text="出库表")

        # ——— 库存表分页 ———
        inventory_frame = ttk.Frame(notebook)
        inventory_frame.pack(fill="both", expand=True)

        # 下方是否隐藏零库存
        self.hide_zero_var = tk.BooleanVar(value=False)  # 修复：默认不隐藏
        hide_zero_cb = ttk.Checkbutton(inventory_frame,
                                       text="是否隐藏数量为0的物资",
                                       variable=self.hide_zero_var,
                                       command=self.refresh_inventory)
        hide_zero_cb.pack(fill="x", pady=5)

        inventory_table_frame = ttk.Frame(inventory_frame)
        inventory_table_frame.pack(fill="both", expand=True)
        inventory_table = InventoryTable(inventory_table_frame, "库存表", inventory_cols)
        inventory_table.pack(fill="both", expand=True)
        inventory_table["show"] = "headings"
        self.inventory_table = inventory_table
        self.tables["库存表"] = inventory_table
        notebook.add(inventory_frame, text="库存表")

        # 底部操作栏
        bottom_frame = ttk.Frame(root)
        bottom_frame.pack(fill="x", pady=5)
        refresh_btn = ttk.Button(bottom_frame, text="刷新所有表", command=self.refresh_all)
        refresh_btn.pack(side="left", padx=5)
        in_btn = ttk.Button(bottom_frame, text="入库", command=self.open_inbound_system)
        in_btn.pack(side="left", padx=5)
        out_btn = ttk.Button(bottom_frame, text="出库", command=self.open_outbound_system)
        out_btn.pack(side="left", padx=5)

        self.refresh_all()

    def refresh_all(self):
        """刷新所有表格"""
        for table in self.tables.values():
            if table.table_name == "库存表":
                table.hide_zero = self.hide_zero_var.get()
            table.load_data(query_filter=self.search_var.get().strip())
        messagebox.showinfo("提示", "所有表已重新加载！")

    def refresh_inventory(self):
        """只刷新库存表的零库存开关后，重新加载数据"""
        self.inventory_table.hide_zero = self.hide_zero_var.get()
        self.refresh_all()

    # —————————— 导出入库表(仍用pandas) ——————————
    def export_inbound(self):
        table = self.tables["入库表"]
        selected = table.selection()
        if not selected:
            messagebox.showwarning("警告","请在入库表中至少选一行")
            return

        rows = []
        for iid in selected:
            vals = table.item(iid,"values")
            row_dict = {}
            # 跳过“序号”列
            for col in self.inbound_cols[1:]:
                try:
                    idx = self.inbound_cols.index(col)
                    row_dict[col] = vals[idx]
                except:
                    continue
            rows.append(row_dict)

        if not rows:
            messagebox.showwarning("警告","无可导出的数据")
            return

        df = pd.DataFrame(rows)
        prefix = rows[0].get("入库日期", datetime.date.today().strftime("%Y-%m-%d")).replace("/","-")
        fname = f"{prefix}入库表.xlsx"
        folder = datetime.date.today().strftime("%Y-%m")
        if not os.path.exists(folder):
            os.makedirs(folder)
        fpath = os.path.join(folder,fname)
        try:
            df.to_excel(fpath, index=False)
            messagebox.showinfo("导出成功", f"入库表已导出为: {fpath}")
        except Exception as e:
            messagebox.showerror("错误",f"导出入库表失败: {e}")

    # —————————— 使用 Excel COM 导出出库领料单(保留页眉图片) ——————————
    def export_outbound(self):
        table = self.tables["出库表"]
        selected = table.selection()
        if not selected:
            messagebox.showwarning("警告","请在出库表中至少选一行")
            return

        data = []
        for iid in selected:
            vals = table.item(iid,"values")
            # 0=序号,1=出库日期,2=物资名称,3=规格型号,4=厂家,5=项目名称,6=项目编码,7=单位,8=数量
            data.append({
                "物资名称": vals[2],
                "规格型号": vals[3],
                "厂家":     vals[4],
                "单位":     vals[7],
                "数量":     vals[8],
            })

        if not data:
            messagebox.showwarning("警告","选中行中无可导出的数据")
            return

        try:
            template_name = "出库领料单.xlsx"
            base_dir = os.path.dirname(os.path.abspath(__file__))
            template_path = os.path.join(base_dir, template_name)
            if not os.path.exists(template_path):
                messagebox.showerror("错误", f"模板文件不存在: {template_path}")
                return

            folder = datetime.date.today().strftime("%Y-%m")
            if not os.path.exists(folder):
                os.makedirs(folder)

            date_str = datetime.date.today().strftime("%Y%m%d")
            new_name = f"{date_str}出库领料单.xlsx"
            new_path = os.path.join(folder,new_name)

            shutil.copy(template_path, new_path)

            # 将路径转换为绝对路径+反斜杠，避免Excel找不到文件
            abspath = os.path.abspath(new_path).replace('/','\\')

            excel = win32.Dispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False

            wb = excel.Workbooks.Open(abspath)
            ws = wb.ActiveSheet

            # 默认从第12行开始写
            start_row = 12
            for i,rowdata in enumerate(data):
                r = start_row + i
                ws.Range(f"B{r}").Value = rowdata["物资名称"]
                ws.Range(f"C{r}").Value = rowdata["规格型号"]
                ws.Range(f"D{r}").Value = rowdata["厂家"]
                ws.Range(f"E{r}").Value = rowdata["单位"]
                ws.Range(f"F{r}").Value = rowdata["数量"]

            wb.Save()
            wb.Close()
            excel.Quit()

            messagebox.showinfo("导出成功", f"出库领料单已生成: {new_path}")

        except Exception as e:
            messagebox.showerror("错误",f"导出出库领料单失败: {e}")

    # 仓库管理界面.py  —— 只替换 generate_weekly_report 函数即可
    # —————————— 生成项目周报(同时导出库存表) ——————————
    def generate_weekly_report(self):
        """导出入库、出库、库存三份数据；库存表不包含“序号”列。"""
        today = datetime.date.today()
        # 计算当周起止
        start = today - datetime.timedelta(days=today.weekday())   # 周一
        end   = start + datetime.timedelta(days=6)                 # 周日
        week  = (today.day - 1) // 7 + 1                           # 月内第几周

        # 绝对基准目录（脚本同级）
        base_dir = os.path.dirname(os.path.abspath(__file__))
        month_folder = datetime.date.today().strftime("%Y-%m")     # 2025-04
        out_dir = os.path.join(base_dir, month_folder)
        os.makedirs(out_dir, exist_ok=True)

        fname = f"{today.year}年{today.month}月第{week}周项目周报.xlsx"
        path  = os.path.join(out_dir, fname)

        conn = connect_db()

        # -------- SQL ---------
        q_in = f"""
            SELECT 入库日期, 物资名称, 规格型号, 专业类别, 单位, 数量, 批次或卷号, 厂家,
                   领料人, 验收人, 委托人, 备注, 项目名称, 项目编码,
                   物资白单单号, 物资属性, 明细账存货编码, 物资黄单单号
            FROM 入库表
            WHERE 入库日期 BETWEEN '{start:%Y/%m/%d}' AND '{end:%Y/%m/%d}'
        """
        q_out = f"""
            SELECT 出库日期, 物资名称, 规格型号, 厂家, 项目名称, 项目编码, 单位, 数量,
                   批次或卷号, 领用人, 经办人, 备注
            FROM 出库表
            WHERE 出库日期 BETWEEN '{start:%Y/%m/%d}' AND '{end:%Y/%m/%d}'
        """
        q_inv = """
            SELECT 物资名称, 规格型号, 厂家, 专业类别, 委托人, 单位,
                   当前库存, 最近入库日期, 最近出库日期, 存放位置, 备注
            FROM 库存表
        """

        try:
            df_in  = pd.read_sql_query(q_in,  conn)
            df_out = pd.read_sql_query(q_out, conn)
            df_inv = pd.read_sql_query(q_inv, conn)
            conn.close()

            # ------ 如果三表都空，直接提示 ------
            if df_in.empty and df_out.empty and df_inv.empty:
                messagebox.showinfo("提示", "本周无入库/出库/库存记录可导出。")
                return

            with pd.ExcelWriter(path, engine='openpyxl') as writer:
                if not df_in.empty:
                    df_in.to_excel(writer,  sheet_name="入库记录", index=False)
                if not df_out.empty:
                    df_out.to_excel(writer, sheet_name="出库记录", index=False)
                if not df_inv.empty:
                    df_inv.to_excel(writer, sheet_name="库存表", index=False)

            # 成功提示 + 打开资源管理器并选中文件，方便你立刻查看
            messagebox.showinfo("生成成功", f"项目周报已生成为：\n{path}")
            try:
                subprocess.Popen(['explorer', '/select,', path])
            except Exception:
                pass

        except Exception as e:
            messagebox.showerror("错误", f"生成项目周报失败：{e}")


    def open_inbound_system(self):
        try:
            base_dir = os.path.dirname(os.path.abspath(__file__))
            script = os.path.join(base_dir, "入库管理系统.py")
            subprocess.Popen(["python", script])
        except Exception as e:
            messagebox.showerror("错误",f"无法启动入库管理系统: {e}")

    def open_outbound_system(self):
        try:
            base_dir = os.path.dirname(os.path.abspath(__file__))
            script = os.path.join(base_dir, "出库管理系统.py")
            subprocess.Popen(["python", script])
        except Exception as e:
            messagebox.showerror("错误",f"无法启动出库管理系统: {e}")

if __name__=="__main__":
    root = tk.Tk()
    app = InventoryApp(root)
    root.mainloop()
